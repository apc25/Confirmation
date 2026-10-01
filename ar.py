"""채권채무조회서: read the counterparty input workbook and generate letters (HWP + PDF)."""
import html
import re
import tempfile
from pathlib import Path

import openpyxl

import generate
from generate import FIELDS, VERSION, InputError, korean_date, merge_pdfs, safe
from hwpauto import Hwp

AR_SHEET = "채권채무"
AR_COLUMNS = [("조회서번호", "선택", "거래처마다 같은 번호. 비우면 생성 순번(001, 002 …)을 씁니다."),
              ("거래처명", "필수", "같은 거래처의 여러 줄은 조회서 한 통으로 묶입니다."),
              ("구분", "필수", "채권 = 당사가 받을 금액(조회서 1번 표), 채무 = 당사가 지급할 금액(2번 표)"),
              ("계정과목", "필수", "예: 외상매출금, 미수금, 외상매입금, 미지급금"),
              ("적요", "선택", ""),
              ("금액", "필수", "숫자"),
              ("비고", "선택", "")]
AR_FIELDS = ["감사인명", "담당CPA", "CPA전화", "회사명", "대표자", "작성일", "조회기준일"]
AR_KINDS = {"채권": "당사가 받아야 할 금액", "채무": "당사가 지급해야 할 금액"}
AR_ROWS = 3
AR_COPIES = ("감사인보관용", "회신용", "회사보관용")


def write_ar_input_template(path):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = AR_SHEET
    ws.append([c for c, _, _ in AR_COLUMNS])
    for cell, width in zip(ws[1], (12, 28, 8, 20, 30, 16, 20)):
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"
    dv = DataValidation(type="list", formula1='"채권,채무"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("C2:C2000")
    ws.column_dimensions["F"].number_format = "#,##0"
    guide = wb.create_sheet("작성방법")
    guide.append(["열", "필수", "설명"])
    for row in AR_COLUMNS:
        guide.append(list(row))
    guide.append([])
    guide.append(["", "", "한 줄에 계정 하나씩 적습니다. 열 순서는 바꿔도 됩니다(1행 제목으로 열을 찾습니다)."])
    guide.append(["", "", f"한 거래처의 채권 또는 채무가 {AR_ROWS}줄을 넘으면 표에는 {AR_ROWS - 1}줄과 "
                          "'기타(별지 참조)'를 넣고, 전체 내역을 별지로 붙입니다."])
    for col, width in zip("ABC", (12, 8, 90)):
        guide.column_dimensions[col].width = width
    wb.save(path)


def parse_amount(v):
    if isinstance(v, (int, float)):
        return v
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def fmt_amount(v):
    return f"{int(v):,}" if float(v).is_integer() else f"{v:,.2f}"


def load_ar(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[AR_SHEET] if AR_SHEET in wb.sheetnames else wb.worksheets[0]
    head = {str(c.value).strip(): c.column for c in ws[1] if c.value}
    missing = [c for c, need, _ in AR_COLUMNS if need == "필수" and c not in head]
    if missing:
        raise InputError(f"채권채무 입력 1행에 필수 열이 없음: {', '.join(missing)}")
    groups, errors = {}, []
    for r in range(2, ws.max_row + 1):
        v = {c: ws.cell(r, head[c]).value if c in head else None for c, _, _ in AR_COLUMNS}
        if all(x in (None, "") for x in v.values()):
            continue
        amount = parse_amount(v["금액"]) if v["금액"] not in (None, "") else None
        problems = [c for c in ("거래처명", "계정과목") if not str(v[c] or "").strip()]
        if v["구분"] not in AR_KINDS:
            problems.append("구분(채권/채무)")
        if amount is None:
            problems.append("금액")
        if problems:
            errors.append(f"{r}행 {', '.join(problems)}")
            continue
        key = (str(v["조회서번호"] or "").strip(), str(v["거래처명"]).strip())
        groups.setdefault(key, []).append({"row": r, "kind": v["구분"], "acct": str(v["계정과목"]).strip(),
                                           "memo": str(v["적요"] or "").strip(), "amount": amount,
                                           "note": str(v["비고"] or "").strip()})
    if errors:
        more = f" 외 {len(errors) - 10}건" if len(errors) > 10 else ""
        raise InputError("채권채무 입력 확인 필요: " + "; ".join(errors[:10]) + more)
    if not groups:
        raise InputError("채권채무 입력 파일에 거래처 줄이 없습니다.")
    return list(groups.items())


def ar_section(sec, lines):
    """Tokens for one amount table; more than AR_ROWS lines -> last row becomes '기타 (별지 참조)'."""
    if not lines:
        rows, total = [("해당사항 없음", "", "", "")], "0"
    else:
        shown = lines if len(lines) <= AR_ROWS else lines[:AR_ROWS - 1]
        rows = [(x["acct"], x["memo"], fmt_amount(x["amount"]), x["note"]) for x in shown]
        if len(lines) > AR_ROWS:
            rest = lines[AR_ROWS - 1:]
            rows.append(("기타", f"별지 참조 ({len(rest)}건)", fmt_amount(sum(x["amount"] for x in rest)), ""))
        total = fmt_amount(sum(x["amount"] for x in lines))
    tokens = {f"{sec}합계": total}
    for i in range(1, AR_ROWS + 1):
        row = rows[i - 1] if i <= len(rows) else ("", "", "", "")
        for col, val in zip(("계정", "적요", "금액", "비고"), row):
            tokens[f"{sec}{i}_{col}"] = val
    return tokens


def ar_detail_html(values, no, name, lines, copy):
    e = html.escape
    body = ""
    for sec, title in AR_KINDS.items():
        sec_lines = [x for x in lines if x["kind"] == sec]
        if not sec_lines:
            continue
        rows = "".join(f'<tr><td width="150">{e(x["acct"])}</td><td width="200">{e(x["memo"])}</td>'
                       f'<td width="130" align="right">{fmt_amount(x["amount"])}</td>'
                       f'<td width="140">{e(x["note"])}</td></tr>' for x in sec_lines)
        body += (f'<p align="left"><b>{sec} — {title}</b></p>'
                 '<table border="1" cellspacing="0" cellpadding="4" width="620">'
                 '<tr><th width="150">계정과목</th><th width="200">적요</th><th width="130">금액</th>'
                 f'<th width="140">비고</th></tr>{rows}'
                 f'<tr><td colspan="2" align="center">합 계</td>'
                 f'<td align="right">{fmt_amount(sum(x["amount"] for x in sec_lines))}</td><td></td></tr></table>')
    return ('<html><head><meta charset="utf-8"></head><body>'
            f'<p align="center"><b><font size="5">채권·채무 명세 (별지)</font></b></p>'
            f'<p align="right">{copy}</p>'
            f'<p align="left">거래처 : {e(name)}&nbsp;&nbsp;&nbsp;조회서번호 : {e(no)}&nbsp;&nbsp;&nbsp;'
            f'조회기준일 : {e(values["조회기준일"])} 현재&nbsp;&nbsp;&nbsp;회사 : {e(values["회사명"])}</p>'
            f'{body}</body></html>')


def prepare_ar(raw):
    labels = {k: lab for group in FIELDS.values() for k, lab in group}
    missing = [labels[k] for k in AR_FIELDS if not str(raw.get(k, "")).strip()]
    if missing:
        raise InputError(f"기본정보 탭에서 비어 있는 항목: {', '.join(missing)}")
    values = {k: str(raw[k]).strip() for k in AR_FIELDS}
    for k in ("작성일", "조회기준일"):
        values[k] = korean_date(values[k], labels[k])
    return values


def run_ar(raw, path, log=print):
    values = prepare_ar(raw)
    letters = load_ar(path)
    for k, v in values.items():
        if "XXX" in v:
            log(f"경고: '{k}' 값이 아직 XXX 자리표시: {v}")
    # Output/template folders are read from generate at call time (exe location, test overrides).
    out = generate.HERE / "out" / safe(f"{values['회사명']}_{re.sub(r'\D', '', values['조회기준일'])}")
    folder = out / "채권채무"
    made = []
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp, Hwp() as h:
        for i, ((no, name), lines) in enumerate(letters, 1):
            no = no or f"{i:03d}"
            secs = {sec: [x for x in lines if x["kind"] == sec] for sec in AR_KINDS}
            overflow = any(len(x) > AR_ROWS for x in secs.values())
            details = []
            if overflow:
                for j, copy in enumerate(AR_COPIES):
                    src = Path(tmp) / f"detail{j}.html"
                    src.write_text(ar_detail_html(values, no, name, lines, copy), encoding="utf-8")
                    h.open(src, "HTML")
                    details.append(Path(tmp) / f"detail{j}.hwp")
                    h.save(details[-1], "HWP")
                    h.close_doc()
            h.open(generate.TPL_DIR / "채권채무.hwp")
            pages = h.app.PageCount
            tokens = {**values, "거래처명": name, **ar_section("채권", secs["채권"]), **ar_section("채무", secs["채무"])}
            for token, v in tokens.items():
                h.replace("{{" + token + "}}", v)
            h.replace_everywhere("{{조회서번호}}", no)
            if h.app.PageCount != pages:
                log(f"경고: {name} — 값을 채운 뒤 {pages}쪽이 {h.app.PageCount}쪽으로 늘어남(긴 값으로 줄이 넘침). PDF 확인 필요")
            for d in details:
                h.append(d)
            left = sorted(set(re.findall(r"\{\{[^}<]*\}\}", h.full_text())))
            if left:
                raise InputError(f"{name}: 채워지지 않은 칸 {left}")
            fname = safe(f"{i:03d}_{name}")
            h.save(folder / f"{fname}.hwp", "HWP")
            h.save(folder / f"{fname}.pdf", "PDF")
            h.close_doc()
            ar_total = {sec: fmt_amount(sum(x["amount"] for x in secs[sec])) for sec in AR_KINDS}
            made.append([i, no, name, ar_total["채권"], ar_total["채무"], len(lines), "별지" if overflow else "",
                         f"{fname}.pdf", ", ".join(str(x["row"]) for x in lines)])
            log(f"채권채무/{fname}.pdf  (조회서번호 {no}, {len(lines)}줄{', 별지 첨부' if overflow else ''})")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "입력값"
    ws.append(["항목", "값"])
    ws.append(["프로그램 버전", VERSION])
    labels = {k: lab for group in FIELDS.values() for k, lab in group}
    for k in AR_FIELDS:
        ws.append([labels[k], raw[k]])
    ws.append(["거래처 입력 파일", str(path)])
    ws2 = wb.create_sheet("생성목록")
    ws2.append(["순번", "조회서번호", "거래처명", "채권 합계", "채무 합계", "줄 수", "별지", "파일", "입력 행"])
    for row in made:
        ws2.append(row)
    for sheet, widths in ((ws, [18, 80]), (ws2, [6, 12, 28, 16, 16, 8, 8, 36, 20])):
        for col, w in zip("ABCDEFGHI", widths):
            sheet.column_dimensions[col].width = w
    wb.save(folder / "생성기록_채권채무.xlsx")
    pdfs = [folder / row[7] for row in made]
    log(f"합본: {merge_pdfs(pdfs, folder / '합본_채권채무.pdf')}쪽 ({len(pdfs)}통) → 합본_채권채무.pdf")
    log(f"출력 폴더: {folder}")
    return folder
