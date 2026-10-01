"""Generate confirmation letters (HWP + PDF) per institution from the control-sheet list.

Input values come from 입력값.json (written by app.py).
Usage: python generate.py ["조회처 목록.xlsx"]
"""
import calendar
import html
import json
import re
import sys
import tempfile
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl

from hwpauto import Hwp

HERE = Path(sys.executable).parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent
TPL_DIR = HERE / "templates"
SAVED = HERE / "입력값.json"
ONLINE_DIR = "전자조회 가능(비상용)"
MAIL_DIR = "우편 발송"
LIST_HEADERS = {2: "용도", 4: "금융기관명", 5: "지점명", 7: "온라인"}

# Common inputs: (token, label), all required.
FIELDS = {
    "감사인": [("감사인명", "감사인명"), ("감사인사업자번호", "사업자등록번호"), ("담당CPA", "담당 공인회계사"),
             ("CPA전화", "전화"), ("CPA팩스", "팩스"), ("CPA이메일", "이메일"),
             ("환급계좌", "수수료 환급계좌"), ("환급은행", "환급 은행"), ("환급예금주", "예금주")],
    "회사": [("회사명", "회사명"), ("법인등록번호", "법인등록번호"), ("대표자", "대표자")],
    "날짜": [("작성일", "작성일"), ("조회기준일", "조회기준일"), ("요구서유효기간", "요구서 유효기간")],
}
SEAL = ("첨부여부", "법인인감증명서 첨부", "yn")
# 용도 value in the control sheet -> tab/template name and form-only inputs (token, label, kind).
# Form-only inputs are stored as "<tab>.<token>" and required only when that form has letters.
FORMS = {
    "은행": {"tab": "은행", "fields": [("교부시작일", "어음·수표 교부기간 시작", "date"),
                                    ("교부종료일", "어음·수표 교부기간 종료", "date"),
                                    ("당좌시작일", "당좌거래 기간 시작", "date"),
                                    ("당좌종료일", "당좌거래 기간 종료", "date")],
             "note": "6번 어음·수표 교부 명세와 10번 당좌거래명세의 기간입니다."},
    "비은행금융회사": {"tab": "비은행", "fields": [], "note": "공통 항목만 씁니다."},
    "여신전문금융회사": {"tab": "여신", "fields": [SEAL], "note": "요구서 붙임의 '(첨부여부 : 여 / 부)'에 표시됩니다."},
    "보험회사": {"tab": "보험", "fields": [SEAL], "note": "요구서 붙임의 '(첨부여부 : 여 / 부)'에 표시됩니다."},
    "증권회사": {"tab": "증권", "fields": [("거래시작일", "거래내역 기간 시작", "date"),
                                    ("거래종료일", "거래내역 기간 종료", "date")],
             "note": "9번 거래내역 첨부 기간입니다."},
    "보증기관": {"tab": "보증기관", "fields": [SEAL], "note": "요구서 붙임의 '(첨부여부 : 여 / 부)'에 표시됩니다."},
}
SEAL_TEXT = {"여": "여 ○  부 __", "부": "여 __  부 ○"}


class InputError(Exception):
    pass


def parse_date(v, label):
    try:
        return datetime.strptime(str(v).strip(), "%Y-%m-%d").date()
    except ValueError:
        raise InputError(f"'{label}' 날짜 형식 오류: {v!r} (예: 2026-12-31)")


def korean_date(v, label):
    d = parse_date(v, label)
    return f"{d.year}년 {d.month}월 {d.day}일"


def eomonth(d, months):
    y, m = divmod(d.month - 1 + months, 12)
    y, m = d.year + y, m + 1
    return date(y, m, calendar.monthrange(y, m)[1])


def derived_dates(base):
    """Default dates that follow 조회기준일 (same rules as the old Excel macro)."""
    d = {"요구서유효기간": eomonth(base, 3),
         "은행.교부시작일": date(base.year - 1, base.month, base.day) + timedelta(days=1),
         "은행.교부종료일": base,
         "은행.당좌시작일": base - timedelta(days=30),
         "은행.당좌종료일": eomonth(base, 1),
         "증권.거래시작일": base - timedelta(days=30),
         "증권.거래종료일": eomonth(base, 1)}
    return {k: v.isoformat() for k, v in d.items()}


def load_saved():
    return json.loads(SAVED.read_text(encoding="utf-8")) if SAVED.exists() else {}


def save_values(raw):
    SAVED.write_text(json.dumps(raw, ensure_ascii=False, indent=2), encoding="utf-8")


def prepare(raw, kinds):
    """Validate raw input for the forms in `kinds`; return common tokens, per-form tokens, business numbers."""
    labels = {k: lab for group in FIELDS.values() for k, lab in group}
    missing = [labels[k] for k in labels if not str(raw.get(k, "")).strip()]
    for kind in kinds:
        tab = FORMS[kind]["tab"]
        missing += [f"{tab} 탭 {lab}" for tok, lab, _ in FORMS[kind]["fields"]
                    if not str(raw.get(f"{tab}.{tok}", "")).strip()]
    if missing:
        raise InputError(f"비어 있는 항목: {', '.join(missing)}")
    biz = [b for b in raw.get("사업자목록", []) if str(b[0]).strip()]
    if not biz:
        raise InputError("사업자등록번호를 하나 이상 넣으세요.")
    values = {k: str(raw[k]).strip() for k in labels}
    for k, lab in FIELDS["날짜"]:
        values[k] = korean_date(values[k], lab)
    values["사업자등록번호"] = biz[0][0].strip() if len(biz) == 1 else "별도 첨부"
    per_form = {}
    for kind in kinds:
        tab = FORMS[kind]["tab"]
        per_form[kind] = {}
        for tok, lab, typ in FORMS[kind]["fields"]:
            v = str(raw[f"{tab}.{tok}"]).strip()
            per_form[kind][tok] = korean_date(v, f"{tab} 탭 {lab}") if typ == "date" else SEAL_TEXT[v]
    return values, per_form, biz


def load_letters(path):
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    for col, label in LIST_HEADERS.items():
        if label not in str(ws.cell(6, col).value or ""):
            raise InputError(f"조회처 목록 6행 {col}열 제목이 '{label}'이 아님 — 양식 변경 여부 확인")
    groups, skipped = defaultdict(list), defaultdict(list)
    for r in range(8, ws.max_row + 1):
        kind, acct, inst, branch, level, online = (ws.cell(r, c).value for c in (2, 3, 4, 5, 6, 7))
        if not (kind or acct or inst):
            continue
        if not inst:
            skipped["금융기관명 공란"].append(r)
        elif kind not in FORMS:
            skipped[f"용도 '{kind or '공란'}'"].append(r)
        else:
            to = f"{inst} {branch}" if branch else (f"{inst} 본점" if level == "본점" else inst)
            groups[(kind, to.strip())].append({"row": r, "acct": acct, "online": online == "O"})
    letters = sorted(groups.items(), key=lambda g: g[1][0]["row"])
    return letters, skipped


def is_online(rows):
    return all(x["online"] for x in rows)


def safe(name):
    return re.sub(r'[\\/:*?"<>|]', "_", name)


def biz_list_html(values, biz):
    e = html.escape
    rows = "".join(f'<tr><td width="60" align="center">{i}</td><td width="220">{e(n.strip())}</td>'
                   f'<td width="340">{e(str(name or "").strip())}</td></tr>' for i, (n, name) in enumerate(biz, 1))
    return ('<html><head><meta charset="utf-8"></head><body>'
            '<p align="center"><b><font size="5">사업자등록번호 목록</font></b></p>'
            f'<p align="left">회사명 : {e(values["회사명"])}&nbsp;&nbsp;&nbsp;&nbsp;법인등록번호 : {e(values["법인등록번호"])}</p>'
            '<table border="1" cellspacing="0" cellpadding="6" width="620">'
            '<tr><th width="60" align="center">번호</th><th width="220" align="center">사업자등록번호</th>'
            f'<th width="340" align="center">사업장명</th></tr>{rows}</table></body></html>')


def write_record(out, raw, kinds, biz, made):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "입력값"
    ws.append(["구분", "항목", "값"])
    for group, fields in FIELDS.items():
        for k, label in fields:
            ws.append([group, label, raw[k]])
    for kind in kinds:
        tab = FORMS[kind]["tab"]
        for tok, label, _ in FORMS[kind]["fields"]:
            ws.append([f"{tab} 탭", label, raw[f"{tab}.{tok}"]])
    ws2 = wb.create_sheet("사업자등록번호")
    ws2.append(["번호", "사업자등록번호", "사업장명"])
    for i, (n, name) in enumerate(biz, 1):
        ws2.append([i, n, name])
    ws3 = wb.create_sheet("생성목록")
    ws3.append(["순번", "양식", "구분", "수신처", "파일", "조회처 목록 행", "계정과목"])
    for row in made:
        ws3.append(row)
    for sheet, widths in ((ws, [12, 26, 40]), (ws2, [8, 20, 30]), (ws3, [6, 10, 22, 30, 40, 16, 40])):
        for col, w in zip("ABCDEFG", widths):
            sheet.column_dimensions[col].width = w
    wb.save(out / "생성기록.xlsx")


def run(raw, list_path, log=print):
    letters, skipped = load_letters(list_path)
    kinds = [k for k in FORMS if any(kind == k for (kind, _), _ in letters)]
    values, per_form, biz = prepare(raw, kinds)
    for k, v in values.items():
        if "XXX" in v:
            log(f"경고: '{k}' 값이 아직 XXX 자리표시: {v}")
    for reason, rows in skipped.items():
        log(f"제외 {len(rows)}행 — {reason}: {rows}")

    out = HERE / "out" / safe(f"{values['회사명']}_{re.sub(r'\D', '', values['조회기준일'])}")
    out.mkdir(parents=True, exist_ok=True)
    made = []
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp, Hwp() as h:
        biz_doc = None
        if len(biz) > 1:
            src = Path(tmp) / "biz.html"
            src.write_text(biz_list_html(values, biz), encoding="utf-8")
            h.open(src, "HTML")
            biz_doc = Path(tmp) / "biz.hwp"
            h.save(biz_doc, "HWP")
            h.close_doc()
            log(f"사업자등록번호 {len(biz)}개 — 목록 쪽을 각 조회서 끝에 붙임")
        for i, ((kind, to), rows) in enumerate(letters, 1):
            tab = FORMS[kind]["tab"]
            if len({x["online"] for x in rows}) > 1:
                log(f"경고: {to} — 온라인 O와 공란이 섞여 있어 우편 발송으로 분류")
            folder = out / (ONLINE_DIR if is_online(rows) else MAIL_DIR)
            name = safe(f"{i:02d}_{tab}_{to}")
            h.open(TPL_DIR / f"{tab}.hwp")
            pages = h.app.PageCount
            for token, v in {**values, **per_form[kind], "수신처": to}.items():
                h.replace("{{" + token + "}}", v)
            if h.app.PageCount != pages:
                log(f"경고: {name} — 값을 채운 뒤 {pages}쪽이 {h.app.PageCount}쪽으로 늘어남(긴 값으로 줄이 넘침). PDF 확인 필요")
            if biz_doc:
                h.append(biz_doc)
            left = sorted(set(re.findall(r"\{\{[^}]*\}\}", h.text())))
            if left:
                raise InputError(f"{name}: 채워지지 않은 칸 {left}")
            h.save(folder / f"{name}.hwp", "HWP")
            h.save(folder / f"{name}.pdf", "PDF")
            h.close_doc()
            accts = ", ".join(sorted({str(x['acct']) for x in rows}))
            made.append([i, tab, folder.name, to, f"{name}.pdf", ", ".join(str(x["row"]) for x in rows), accts])
            log(f"{folder.name}/{name}.pdf  (목록 {[x['row'] for x in rows]}행, {accts})")
    write_record(out, raw, kinds, biz, made)
    mail = [out / MAIL_DIR / row[4] for row in made if row[2] == MAIL_DIR]
    if mail:
        log(f"합본: {merge_pdfs(mail, out / '합본_우편발송.pdf')}쪽 ({len(mail)}통) → 합본_우편발송.pdf")
    else:
        log("우편 발송 대상이 없어 합본은 만들지 않음")
    log(f"출력 폴더: {out}")
    return out


def merge_pdfs(paths, dst):
    """Concatenate PDFs in the given order; returns the page count. Merged at PDF level so headers stay per letter."""
    from pypdf import PdfWriter
    w = PdfWriter()
    for p in paths:
        w.append(str(p))
    with open(dst, "wb") as f:
        w.write(f)
    return len(w.pages)


# ---- 채권채무조회서 ----
AR_SHEET = "채권채무"
AR_COLUMNS = [("조회서번호", "선택", "거래처마다 같은 번호. 비우면 생성 순번(001, 002 …)을 씁니다."),
              ("거래처명", "필수", "같은 거래처의 여러 줄은 조회서 한 통으로 묶입니다."),
              ("구분", "필수", "채권 = 당사가 받을 금액(조회서 1번 표), 채무 = 당사가 지급할 금액(2번 표)"),
              ("계정과목", "필수", "예: 외상매출금, 미수금, 외상매입금, 미지급금"),
              ("적요", "선택", ""),
              ("금액", "필수", "숫자"),
              ("비고", "선택", "")]
AR_FIELDS = ["감사인명", "담당CPA", "CPA전화", "회사명", "대표자", "작성일", "조회기준일"]
AR_KINDS = {"채권": "당사가 받아야 할 금액", "채무": "당사가 지급해야 할 금액"}
AR_ROWS = 3
AR_COPIES = ("감사인보관용", "회신용", "회사보관용")


def write_ar_input_template(path):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = AR_SHEET
    ws.append([c for c, _, _ in AR_COLUMNS])
    for cell, width in zip(ws[1], (12, 28, 8, 20, 30, 16, 20)):
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"
    dv = DataValidation(type="list", formula1='"채권,채무"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("C2:C2000")
    ws.column_dimensions["F"].number_format = "#,##0"
    guide = wb.create_sheet("작성방법")
    guide.append(["열", "필수", "설명"])
    for row in AR_COLUMNS:
        guide.append(list(row))
    guide.append([])
    guide.append(["", "", "한 줄에 계정 하나씩 적습니다. 열 순서는 바꿔도 됩니다(1행 제목으로 열을 찾습니다)."])
    guide.append(["", "", f"한 거래처의 채권 또는 채무가 {AR_ROWS}줄을 넘으면 표에는 {AR_ROWS - 1}줄과 "
                          "'기타(별지 참조)'를 넣고, 전체 내역을 별지로 붙입니다."])
    for col, width in zip("ABC", (12, 8, 90)):
        guide.column_dimensions[col].width = width
    wb.save(path)


def parse_amount(v):
    if isinstance(v, (int, float)):
        return v
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def fmt_amount(v):
    return f"{int(v):,}" if float(v).is_integer() else f"{v:,.2f}"


def load_ar(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[AR_SHEET] if AR_SHEET in wb.sheetnames else wb.worksheets[0]
    head = {str(c.value).strip(): c.column for c in ws[1] if c.value}
    missing = [c for c, need, _ in AR_COLUMNS if need == "필수" and c not in head]
    if missing:
        raise InputError(f"채권채무 입력 1행에 필수 열이 없음: {', '.join(missing)}")
    groups, errors = {}, []
    for r in range(2, ws.max_row + 1):
        v = {c: ws.cell(r, head[c]).value if c in head else None for c, _, _ in AR_COLUMNS}
        if all(x in (None, "") for x in v.values()):
            continue
        amount = parse_amount(v["금액"]) if v["금액"] not in (None, "") else None
        problems = [c for c in ("거래처명", "계정과목") if not str(v[c] or "").strip()]
        if v["구분"] not in AR_KINDS:
            problems.append("구분(채권/채무)")
        if amount is None:
            problems.append("금액")
        if problems:
            errors.append(f"{r}행 {', '.join(problems)}")
            continue
        key = (str(v["조회서번호"] or "").strip(), str(v["거래처명"]).strip())
        groups.setdefault(key, []).append({"row": r, "kind": v["구분"], "acct": str(v["계정과목"]).strip(),
                                           "memo": str(v["적요"] or "").strip(), "amount": amount,
                                           "note": str(v["비고"] or "").strip()})
    if errors:
        more = f" 외 {len(errors) - 10}건" if len(errors) > 10 else ""
        raise InputError("채권채무 입력 확인 필요: " + "; ".join(errors[:10]) + more)
    if not groups:
        raise InputError("채권채무 입력 파일에 거래처 줄이 없습니다.")
    return list(groups.items())


def ar_section(sec, lines):
    """Tokens for one amount table; more than AR_ROWS lines -> last row becomes '기타 (별지 참조)'."""
    if not lines:
        rows, total = [("해당사항 없음", "", "", "")], "0"
    else:
        shown = lines if len(lines) <= AR_ROWS else lines[:AR_ROWS - 1]
        rows = [(x["acct"], x["memo"], fmt_amount(x["amount"]), x["note"]) for x in shown]
        if len(lines) > AR_ROWS:
            rest = lines[AR_ROWS - 1:]
            rows.append(("기타", f"별지 참조 ({len(rest)}건)", fmt_amount(sum(x["amount"] for x in rest)), ""))
        total = fmt_amount(sum(x["amount"] for x in lines))
    tokens = {f"{sec}합계": total}
    for i in range(1, AR_ROWS + 1):
        row = rows[i - 1] if i <= len(rows) else ("", "", "", "")
        for col, val in zip(("계정", "적요", "금액", "비고"), row):
            tokens[f"{sec}{i}_{col}"] = val
    return tokens


def ar_detail_html(values, no, name, lines, copy):
    e = html.escape
    body = ""
    for sec, title in AR_KINDS.items():
        sec_lines = [x for x in lines if x["kind"] == sec]
        if not sec_lines:
            continue
        rows = "".join(f'<tr><td width="150">{e(x["acct"])}</td><td width="200">{e(x["memo"])}</td>'
                       f'<td width="130" align="right">{fmt_amount(x["amount"])}</td>'
                       f'<td width="140">{e(x["note"])}</td></tr>' for x in sec_lines)
        body += (f'<p align="left"><b>{sec} — {title}</b></p>'
                 '<table border="1" cellspacing="0" cellpadding="4" width="620">'
                 '<tr><th width="150">계정과목</th><th width="200">적요</th><th width="130">금액</th>'
                 f'<th width="140">비고</th></tr>{rows}'
                 f'<tr><td colspan="2" align="center">합 계</td>'
                 f'<td align="right">{fmt_amount(sum(x["amount"] for x in sec_lines))}</td><td></td></tr></table>')
    return ('<html><head><meta charset="utf-8"></head><body>'
            f'<p align="center"><b><font size="5">채권·채무 명세 (별지)</font></b></p>'
            f'<p align="right">{copy}</p>'
            f'<p align="left">거래처 : {e(name)}&nbsp;&nbsp;&nbsp;조회서번호 : {e(no)}&nbsp;&nbsp;&nbsp;'
            f'조회기준일 : {e(values["조회기준일"])} 현재&nbsp;&nbsp;&nbsp;회사 : {e(values["회사명"])}</p>'
            f'{body}</body></html>')


def prepare_ar(raw):
    labels = {k: lab for group in FIELDS.values() for k, lab in group}
    missing = [labels[k] for k in AR_FIELDS if not str(raw.get(k, "")).strip()]
    if missing:
        raise InputError(f"기본정보 탭에서 비어 있는 항목: {', '.join(missing)}")
    values = {k: str(raw[k]).strip() for k in AR_FIELDS}
    for k in ("작성일", "조회기준일"):
        values[k] = korean_date(values[k], labels[k])
    return values


def run_ar(raw, path, log=print):
    values = prepare_ar(raw)
    letters = load_ar(path)
    for k, v in values.items():
        if "XXX" in v:
            log(f"경고: '{k}' 값이 아직 XXX 자리표시: {v}")
    out = HERE / "out" / safe(f"{values['회사명']}_{re.sub(r'\D', '', values['조회기준일'])}")
    folder = out / "채권채무"
    made = []
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp, Hwp() as h:
        for i, ((no, name), lines) in enumerate(letters, 1):
            no = no or f"{i:03d}"
            secs = {sec: [x for x in lines if x["kind"] == sec] for sec in AR_KINDS}
            overflow = any(len(x) > AR_ROWS for x in secs.values())
            details = []
            if overflow:
                for j, copy in enumerate(AR_COPIES):
                    src = Path(tmp) / f"detail{j}.html"
                    src.write_text(ar_detail_html(values, no, name, lines, copy), encoding="utf-8")
                    h.open(src, "HTML")
                    details.append(Path(tmp) / f"detail{j}.hwp")
                    h.save(details[-1], "HWP")
                    h.close_doc()
            h.open(TPL_DIR / "채권채무.hwp")
            pages = h.app.PageCount
            tokens = {**values, "거래처명": name, **ar_section("채권", secs["채권"]), **ar_section("채무", secs["채무"])}
            for token, v in tokens.items():
                h.replace("{{" + token + "}}", v)
            h.replace_everywhere("{{조회서번호}}", no)
            if h.app.PageCount != pages:
                log(f"경고: {name} — 값을 채운 뒤 {pages}쪽이 {h.app.PageCount}쪽으로 늘어남(긴 값으로 줄이 넘침). PDF 확인 필요")
            for d in details:
                h.append(d)
            left = sorted(set(re.findall(r"\{\{[^}<]*\}\}", h.full_text())))
            if left:
                raise InputError(f"{name}: 채워지지 않은 칸 {left}")
            fname = safe(f"{i:03d}_{name}")
            h.save(folder / f"{fname}.hwp", "HWP")
            h.save(folder / f"{fname}.pdf", "PDF")
            h.close_doc()
            ar_total = {sec: fmt_amount(sum(x["amount"] for x in secs[sec])) for sec in AR_KINDS}
            made.append([i, no, name, ar_total["채권"], ar_total["채무"], len(lines), "별지" if overflow else "",
                         f"{fname}.pdf", ", ".join(str(x["row"]) for x in lines)])
            log(f"채권채무/{fname}.pdf  (조회서번호 {no}, {len(lines)}줄{', 별지 첨부' if overflow else ''})")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "입력값"
    ws.append(["항목", "값"])
    labels = {k: lab for group in FIELDS.values() for k, lab in group}
    for k in AR_FIELDS:
        ws.append([labels[k], raw[k]])
    ws.append(["거래처 입력 파일", str(path)])
    ws2 = wb.create_sheet("생성목록")
    ws2.append(["순번", "조회서번호", "거래처명", "채권 합계", "채무 합계", "줄 수", "별지", "파일", "입력 행"])
    for row in made:
        ws2.append(row)
    for sheet, widths in ((ws, [18, 80]), (ws2, [6, 12, 28, 16, 16, 8, 8, 36, 20])):
        for col, w in zip("ABCDEFGHI", widths):
            sheet.column_dimensions[col].width = w
    wb.save(folder / "생성기록_채권채무.xlsx")
    pdfs = [folder / row[7] for row in made]
    log(f"합본: {merge_pdfs(pdfs, folder / '합본_채권채무.pdf')}쪽 ({len(pdfs)}통) → 합본_채권채무.pdf")
    log(f"출력 폴더: {folder}")
    return folder


def main():
    try:
        raw = load_saved()
        if len(sys.argv) > 1:
            list_path = Path(sys.argv[1])
        elif raw.get("조회처목록"):
            list_path = Path(raw["조회처목록"])
        else:
            raise InputError("조회처 목록 파일 경로를 인자로 주세요.")
        run(raw, list_path)
    except InputError as e:
        sys.exit(str(e))


if __name__ == "__main__":
    main()
